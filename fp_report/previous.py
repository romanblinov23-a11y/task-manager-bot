"""Чтение отчёта за прошлый месяц — из него берём то, что управляющий уже писал и что нельзя вытащить из НИМБ.

Из прошлого отчёта:
- ФИО управляющего (шапка);
- задачи на следующий месяц (раздел 8) — это задачи прошлого месяца для текущего отчёта (раздел 7);
- факт операционных показателей (раздел 6) — это «прошлый месяц» для текущего отчёта;
- план выплат на следующий месяц (раздел 5) — сверяем с фактом выплат.

Отличаем шаблон от отчёта: у шаблона в шапке пустой «Отчётный месяц».
"""

import docx

from fp_report.build import ASK, _MONTH_RU, MANUAL_OPS


def _cell_text(cell) -> str:
    return cell.text.strip()


def is_blank_template(path: str) -> bool:
    doc = docx.Document(path)
    return not _cell_text(doc.tables[0].rows[1].cells[1])


def month_label_of(path: str) -> str:
    """Отчётный месяц из шапки отчёта, например «Сентябрь 2026»."""
    doc = docx.Document(path)
    return _cell_text(doc.tables[0].rows[1].cells[1])


def _non_empty_lines(table, col: int = 0) -> list[str]:
    out = []
    for row in table.rows[1:]:
        text = _cell_text(row.cells[col])
        if text and text != ASK:
            out.append(text)
    return out


def read_previous(path: str) -> dict:
    doc = docx.Document(path)
    tables = doc.tables

    ops_fact = {}
    for row in tables[10].rows[1:]:
        label = _cell_text(row.cells[0])
        if label in MANUAL_OPS:
            fact = _cell_text(row.cells[2])
            if fact and fact != ASK:
                ops_fact[label] = fact

    return {
        "month_label": _cell_text(tables[0].rows[1].cells[1]),
        "manager_name": _cell_text(tables[0].rows[2].cells[1]),
        "tasks_next": _non_empty_lines(tables[12]),
        "payouts_plan": _non_empty_lines(tables[9]),
        "ops_fact": ops_fact,
    }


def expected_label(month: str) -> str:
    y, m = map(int, month.split("-"))
    return f"{_MONTH_RU[m]} {y}"
